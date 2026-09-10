"""
Module 1 — Automated WHO GLASS AMR data ingestion.

Two cooperating agents:

  scrape_glass()  – headless-browser agent that opens the GLASS dashboard,
                    applies Region / Infection type / Pathogen / Antibiotic
                    filters and captures the exported dataset (PDF/CSV/XLSX).
  etl_load()      – processing agent that validates, cleans, de-duplicates and
                    version-loads raw files into `glass_amr_observations`.

run() orchestrates both, deposits raw files into the timestamp-partitioned raw
bucket (gs://amr-raw-data/glass/YYYY/MM/) and records a `glass_ingestion_runs`
row. Designed to be invoked monthly (Cloud Scheduler → Cloud Run job) or
on-demand from the admin panel.

CLI:  python -m backend.jobs.glass_ingest --region "South-East Asia Region" \
        --infection BLOOD --pathogen "Escherichia coli" --antibiotic "Carbapenems"
"""
from __future__ import annotations

import argparse
import csv
import hashlib
import io
import json
import re
import tempfile
import uuid
from datetime import datetime, timezone
from pathlib import Path
from typing import Any, Dict, List, Optional

from google.cloud import bigquery

from backend.core import storage
from backend.database.bigquery_client import BigQueryClient

GLASS_URL = (
    "https://worldhealthorg.shinyapps.io/glass-dashboard/"
    "_w_2fab8dfb/_w_8a5b788ee39e4590a20d555274dd3e48/#!/amr"
)
RAW_ROOT = "glass"


# ---------------------------------------------------------------------------
# Agent 1 — headless browser scraper
# ---------------------------------------------------------------------------
def scrape_glass(
    query: Dict[str, Optional[str]],
    dest_dir: Path,
    timeout_ms: int = 90_000,
) -> Dict[str, Any]:
    """Best-effort. Returns {files: [...], diagnostics: {...}}.

    The GLASS portal is an R/Shiny app; selectors change over time. Every UI
    step is defensive — if a control is not found we continue and still try to
    capture whatever export / table is available.
    """
    try:
        from playwright.sync_api import sync_playwright
    except Exception as exc:  # noqa: BLE001
        return {"files": [], "diagnostics": {"error": f"playwright unavailable: {exc}"}}

    dest_dir.mkdir(parents=True, exist_ok=True)
    files: List[Path] = []
    diag: Dict[str, Any] = {"steps": [], "url": GLASS_URL}

    def _try(label, fn):
        try:
            fn()
            diag["steps"].append(f"ok: {label}")
        except Exception as e:  # noqa: BLE001
            diag["steps"].append(f"skip: {label} ({type(e).__name__})")

    with sync_playwright() as p:
        browser = p.chromium.launch(headless=True)
        ctx = browser.new_context(accept_downloads=True)
        page = ctx.new_page()
        page.set_default_timeout(timeout_ms)
        try:
            page.goto(GLASS_URL, wait_until="load")
            _try("wait for shiny idle", lambda: page.wait_for_load_state("networkidle"))
            page.wait_for_timeout(6000)

            # Apply filters — Shiny selectize widgets are labelled controls.
            filter_map = [
                ("Region", query.get("region")),
                ("Infection", query.get("infection_type")),
                ("Bacterial pathogen", query.get("pathogen")),
                ("Pathogen", query.get("pathogen")),
                ("Antibiotic", query.get("antibiotic")),
            ]
            for label, value in filter_map:
                if not value:
                    continue

                def _set(label=label, value=value):
                    loc = page.get_by_label(re.compile(label, re.I)).first
                    loc.click()
                    page.keyboard.type(value, delay=40)
                    page.wait_for_timeout(800)
                    page.keyboard.press("Enter")
                    page.wait_for_timeout(1500)

                _try(f"set filter {label}={value}", _set)

            page.wait_for_timeout(4000)

            # Try an explicit download / export control.
            captured = False
            for name in (r"download", r"export", r"\.pdf", r"report", r"csv"):
                try:
                    btn = page.get_by_role("button", name=re.compile(name, re.I)).first
                    if btn.count() == 0:
                        btn = page.get_by_role("link", name=re.compile(name, re.I)).first
                    with page.expect_download(timeout=20_000) as dl_info:
                        btn.click()
                    dl = dl_info.value
                    target = dest_dir / (dl.suggested_filename or f"glass_{uuid.uuid4().hex}.bin")
                    dl.save_as(str(target))
                    files.append(target)
                    diag["steps"].append(f"ok: downloaded {target.name}")
                    captured = True
                    break
                except Exception:  # noqa: BLE001
                    continue

            # Fallback — screenshot + scrape any rendered HTML tables to CSV.
            shot = dest_dir / "glass_dashboard.png"
            _try("screenshot", lambda: page.screenshot(path=str(shot), full_page=True))
            if shot.exists():
                files.append(shot)

            if not captured:
                tables = page.query_selector_all("table")
                for i, tbl in enumerate(tables):
                    rows = []
                    for tr in tbl.query_selector_all("tr"):
                        cells = [c.inner_text().strip() for c in tr.query_selector_all("th,td")]
                        if cells:
                            rows.append(cells)
                    if len(rows) > 1:
                        out = dest_dir / f"glass_table_{i}.csv"
                        with out.open("w", newline="") as fh:
                            csv.writer(fh).writerows(rows)
                        files.append(out)
                        diag["steps"].append(f"ok: extracted table {i} ({len(rows)} rows)")
        finally:
            ctx.close()
            browser.close()

    diag["files_captured"] = [f.name for f in files]
    return {"files": files, "diagnostics": diag}


# ---------------------------------------------------------------------------
# Agent 2 — validate / clean / dedupe / version-load
# ---------------------------------------------------------------------------
_COLUMN_ALIASES = {
    "country_code": ["iso3", "countrycode", "country_code"],
    "country": ["countryterritoryarea", "country", "countryname"],
    "who_region": ["whoregionname", "whoregion", "region"],
    "year": ["year"],
    "specimen": ["specimen", "infectiontype", "infection_type"],
    "pathogen": ["pathogenname", "pathogen", "bacteria"],
    "antibiotic": ["abtargets", "antibiotic", "antibioticclass"],
    "total_specimen_isolates": ["totalspecimenisolates", "totalisolates", "total"],
    "interpretable_ast": ["interpretableast", "interpretable"],
    "resistant_count": ["resistant", "resistantcount"],
    "resistance_rate": ["percentresistant", "resistancerate", "pctresistant"],
}


def _norm(s: str) -> str:
    return re.sub(r"[^a-z0-9]", "", s.lower())


def _read_tabular(name: str, data: bytes) -> List[Dict[str, str]]:
    lname = name.lower()
    if lname.endswith((".xlsx", ".xls")):
        try:
            from openpyxl import load_workbook

            wb = load_workbook(io.BytesIO(data), read_only=True, data_only=True)
            ws = wb.active
            rows = list(ws.iter_rows(values_only=True))
            if not rows:
                return []
            header = [str(h or "") for h in rows[0]]
            return [dict(zip(header, [("" if c is None else str(c)) for c in r])) for r in rows[1:]]
        except Exception:  # noqa: BLE001
            return []
    if lname.endswith(".csv"):
        text = data.decode("utf-8", errors="ignore")
        lines = text.splitlines()
        # GLASS exports bundle several blocks separated by blank lines, each with
        # its own header. Pick the richest block: one that has year + isolate
        # counts + country (the "individual CTA" block), else the first block
        # that has year.
        best_start, best_score = None, -1
        for i, ln in enumerate(lines):
            n = _norm(ln)
            if "year" not in n or ln.count(",") < 2:
                continue
            score = sum(tok in n for tok in
                        ("totalspecimenisolates", "interpretableast", "resistant", "iso3", "percentresistant"))
            if score > best_score:
                best_score, best_start = score, i
            if score >= 4:
                break
        if best_start is None:
            return []
        block = []
        for ln in lines[best_start:]:
            if ln.strip() == "" and block:
                break
            block.append(ln)
        reader = csv.DictReader(block)
        rows: List[Dict[str, str]] = []
        for r in reader:
            rows.append({k: v for k, v in r.items() if k is not None})
        return rows
    return []


def _map_row(raw: Dict[str, str]) -> Optional[Dict[str, Any]]:
    keyed = {_norm(k): v for k, v in raw.items() if k is not None}
    out: Dict[str, Any] = {}
    for target, aliases in _COLUMN_ALIASES.items():
        val = None
        for a in aliases:
            if _norm(a) in keyed and str(keyed[_norm(a)]).strip() not in ("", "nan", "None"):
                val = keyed[_norm(a)]
                break
        out[target] = val

    if not out.get("year"):
        return None
    try:
        out["year"] = int(float(out["year"]))
    except (ValueError, TypeError):
        return None

    for c in ("country_code", "country", "who_region", "specimen", "pathogen", "antibiotic"):
        out[c] = (str(out[c]).strip() if out[c] else "unknown")
    for c in ("total_specimen_isolates", "interpretable_ast", "resistant_count"):
        try:
            out[c] = int(float(out[c])) if out[c] not in (None, "") else 0
        except (ValueError, TypeError):
            out[c] = 0
    try:
        out["resistance_rate"] = float(out["resistance_rate"]) if out["resistance_rate"] not in (None, "") else 0.0
    except (ValueError, TypeError):
        out["resistance_rate"] = 0.0
    if out["resistance_rate"] == 0.0 and out["interpretable_ast"]:
        out["resistance_rate"] = round(100.0 * out["resistant_count"] / out["interpretable_ast"], 2)
    return out


def _row_hash(r: Dict[str, Any]) -> str:
    key = "|".join(str(r.get(k, "")) for k in
                    ("country_code", "who_region", "year", "specimen", "pathogen", "antibiotic"))
    return hashlib.sha256(key.encode()).hexdigest()


def etl_load(
    raw_objects: List[storage.StoredObject],
    dataset_version: str,
    bq: Optional[BigQueryClient] = None,
) -> Dict[str, Any]:
    bq = bq or BigQueryClient()
    table = f"{bq.project_id}.{bq.dataset_id}.glass_amr_observations"

    existing = {
        r["observation_id"]
        for r in bq.query(f"SELECT observation_id FROM `{table}`")
    }

    raw_be = storage.backend(storage.RAW_BUCKET)
    mapped: List[Dict[str, Any]] = []
    parsed_files = 0
    for obj in raw_objects:
        if not obj.key.lower().endswith((".csv", ".xlsx", ".xls")):
            continue
        data = raw_be.get(obj.key)
        records = _read_tabular(obj.key, data)
        if not records:
            continue
        parsed_files += 1
        for raw in records:
            m = _map_row(raw)
            if m:
                mapped.append(m)

    seen: set[str] = set()
    to_insert: List[Dict[str, Any]] = []
    duplicates = 0
    for m in mapped:
        h = _row_hash(m)
        obs_id = f"GLASS-{h[:20]}"
        if obs_id in existing or obs_id in seen:
            duplicates += 1
            continue
        seen.add(obs_id)
        to_insert.append({
            "observation_id": obs_id,
            "country_code": m["country_code"],
            "country": m["country"],
            "who_region": m["who_region"],
            "year": m["year"],
            "specimen": m["specimen"],
            "pathogen": m["pathogen"],
            "antibiotic": m["antibiotic"],
            "total_specimen_isolates": m["total_specimen_isolates"],
            "interpretable_ast": m["interpretable_ast"],
            "resistant_count": m["resistant_count"],
            "resistance_rate": m["resistance_rate"],
            "source": f"glass_ingest:{dataset_version}",
        })

    inserted = 0
    for i in range(0, len(to_insert), 2000):
        chunk = to_insert[i:i + 2000]
        errs = bq.client.insert_rows_json(table, chunk)
        if errs:
            raise RuntimeError(f"glass load error: {errs[:2]}")
        inserted += len(chunk)

    return {
        "files_parsed": parsed_files,
        "rows_mapped": len(mapped),
        "rows_ingested": inserted,
        "rows_deduplicated": duplicates,
    }


# ---------------------------------------------------------------------------
# Orchestrator
# ---------------------------------------------------------------------------
def run(
    query: Optional[Dict[str, Optional[str]]] = None,
    run_type: str = "manual",
    triggered_by: str = "system",
    do_scrape: bool = True,
    bq: Optional[BigQueryClient] = None,
) -> Dict[str, Any]:
    bq = bq or BigQueryClient()
    mark_stale_runs(bq)
    query = query or {"region": None, "infection_type": None, "pathogen": None, "antibiotic": None}
    run_id = f"GRUN-{datetime.now(timezone.utc):%Y%m%dT%H%M%SZ}-{uuid.uuid4().hex[:6]}"
    started = datetime.now(timezone.utc)
    dataset_version = started.strftime("%Y.%m.%d")
    raw_be = storage.backend(storage.RAW_BUCKET)

    record: Dict[str, Any] = {
        "run_id": run_id,
        "run_type": run_type,
        "status": "running",
        "started_at": started.isoformat(),
        "source_url": GLASS_URL,
        "query_params": json.dumps(query),
        "dataset_version": dataset_version,
        "triggered_by": triggered_by,
        "files_downloaded": 0,
        "rows_ingested": 0,
        "rows_deduplicated": 0,
    }
    _write_run(bq, record)

    raw_uris: List[str] = []
    diagnostics: Dict[str, Any] = {}
    status = "success"
    error = None

    try:
        stored: List[storage.StoredObject] = []
        if do_scrape:
            with tempfile.TemporaryDirectory() as tmp:
                scrape = scrape_glass(query, Path(tmp))
                diagnostics = scrape["diagnostics"]
                for fpath in scrape["files"]:
                    key = storage.timestamp_key(RAW_ROOT, Path(fpath).name, started)
                    obj = raw_be.put(key, Path(fpath).read_bytes())
                    stored.append(obj)
                    raw_uris.append(obj.uri)
        # Also pick up any files already dropped in this month's raw prefix.
        month_objs = raw_be.list(storage.month_prefix(RAW_ROOT, started))
        drop_objs = [o for o in month_objs if o.uri not in raw_uris]
        stored.extend(drop_objs)

        data_files = [o for o in stored if o.key.lower().endswith((".csv", ".xlsx", ".xls"))]
        etl = etl_load(data_files, dataset_version, bq) if data_files else {
            "files_parsed": 0, "rows_mapped": 0, "rows_ingested": 0, "rows_deduplicated": 0,
        }

        tabular_captured = any(
            o.key.lower().endswith((".csv", ".xlsx", ".xls")) for o in stored
        )
        if do_scrape and not raw_uris:
            status = "partial"
            error = "Browser agent captured nothing from the portal (UI may have changed). ETL ran on existing raw drops."
        elif do_scrape and not tabular_captured:
            status = "partial"
            error = ("Browser agent captured page artefacts (screenshot/chart) but no tabular export. "
                     "Drop a GLASS CSV/XLSX into the raw bucket for this month to load rows.")

        record.update({
            "status": status,
            "finished_at": datetime.now(timezone.utc).isoformat(),
            "raw_object_uris": json.dumps(raw_uris + [o.uri for o in drop_objs]),
            "files_downloaded": len(raw_uris),
            "rows_ingested": etl["rows_ingested"],
            "rows_deduplicated": etl["rows_deduplicated"],
            "error": error,
        })
    except Exception as exc:  # noqa: BLE001
        status = "failed"
        record.update({
            "status": "failed",
            "finished_at": datetime.now(timezone.utc).isoformat(),
            "error": str(exc)[:1000],
        })
    _write_run(bq, record, update=True)
    record["diagnostics"] = diagnostics
    return record


def mark_stale_runs(bq: BigQueryClient, older_than_minutes: int = 60) -> int:
    """Fail runs stuck in 'running' (e.g. a worker died mid-job) so the run
    history never shows a run spinning forever. Best-effort; returns row count."""
    table = f"{bq.project_id}.{bq.dataset_id}.glass_ingestion_runs"
    try:
        job = bq.client.query(f"""
            UPDATE `{table}`
            SET status = 'failed',
                finished_at = CURRENT_TIMESTAMP(),
                error = COALESCE(error, 'run did not report completion (marked stale)')
            WHERE status = 'running'
              AND started_at < TIMESTAMP_SUB(CURRENT_TIMESTAMP(), INTERVAL {int(older_than_minutes)} MINUTE)
        """)
        job.result()
        return job.num_dml_affected_rows or 0
    except Exception:  # noqa: BLE001
        return 0


def _write_run(bq: BigQueryClient, record: Dict[str, Any], update: bool = False) -> None:
    table = f"{bq.project_id}.{bq.dataset_id}.glass_ingestion_runs"
    if not update:
        # load job (not streaming) so the row can be UPDATEd when the job finishes
        bq.load_rows("glass_ingestion_runs", [record])
        return
    sets = []
    params = [bigquery.ScalarQueryParameter("run_id", "STRING", record["run_id"])]
    for k in ("status", "finished_at", "raw_object_uris", "error"):
        if record.get(k) is not None:
            sets.append(f"{k} = @{k}")
            params.append(bigquery.ScalarQueryParameter(k, "STRING", str(record[k])))
    for k in ("files_downloaded", "rows_ingested", "rows_deduplicated"):
        sets.append(f"{k} = @{k}")
        params.append(bigquery.ScalarQueryParameter(k, "INT64", int(record.get(k, 0))))
    bq.client.query(
        f"UPDATE `{table}` SET {', '.join(sets)} WHERE run_id = @run_id",
        job_config=bigquery.QueryJobConfig(query_parameters=params),
    ).result()


if __name__ == "__main__":
    ap = argparse.ArgumentParser()
    ap.add_argument("--region")
    ap.add_argument("--infection")
    ap.add_argument("--pathogen")
    ap.add_argument("--antibiotic")
    ap.add_argument("--no-scrape", action="store_true")
    a = ap.parse_args()
    out = run(
        query={"region": a.region, "infection_type": a.infection, "pathogen": a.pathogen, "antibiotic": a.antibiotic},
        run_type="manual",
        do_scrape=not a.no_scrape,
    )
    print(json.dumps({k: v for k, v in out.items() if k != "diagnostics"}, indent=2))
    print("diagnostics:", json.dumps(out.get("diagnostics", {}), indent=2)[:2000])
